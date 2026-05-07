from __future__ import annotations

import io
from datetime import datetime

import pandas as pd
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from openpyxl.styles import Font, PatternFill

from typing import cast
from openpyxl.workbook.workbook import Workbook

from utils.models import AnalysisResult
from utils.preflight import evaluate_preflight

def _safe_sheet_name(name: str) -> str:
    invalid = ["\\", "/", "*", "[", "]", ":", "?"]
    for char in invalid:
        name = name.replace(char, "_")
    return name[:31] if name else "Sheet1"


def _write_df(writer: pd.ExcelWriter, df: pd.DataFrame, sheet_name: str) -> None:
    safe_name = _safe_sheet_name(sheet_name)
    if df is None or df.empty:
        pd.DataFrame({"message": ["No data available"]}).to_excel(
            writer, sheet_name=safe_name, index=False
        )
        return
    df.to_excel(writer, sheet_name=safe_name, index=False)


def _build_metadata_sheet(result: AnalysisResult) -> pd.DataFrame:
    return pd.DataFrame(
        [
            {"key": "generated_at_utc", "value": datetime.utcnow().isoformat()},
            {"key": "dataset_type", "value": result.dataset_type},
            {"key": "row_count", "value": len(result.cleaned_df)},
            {"key": "column_count", "value": len(result.cleaned_df.columns)},
            {"key": "alert_count", "value": len(result.alerts) if result.alerts is not None else 0},
            {
                "key": "validation_issue_count",
                "value": len(result.validation_report.issues) if result.validation_report else 0,
            },
            {"key": "join_report_count", "value": len(result.join_reports) if result.join_reports else 0},
        ]
    )


def _build_kpi_sheet(result: AnalysisResult) -> pd.DataFrame:
    if not result.kpis:
        return pd.DataFrame({"Metric": [], "Value": []})
    return pd.DataFrame(list(result.kpis.items()), columns=["Metric", "Value"])


def _build_validation_sheet(result: AnalysisResult) -> pd.DataFrame:
    if result.validation_report is None or not result.validation_report.issues:
        return pd.DataFrame(
            [{"severity": "info", "code": "none", "dataset": result.dataset_type, "message": "No validation issues detected."}]
        )

    rows = []
    for issue in result.validation_report.issues:
        rows.append(
            {
                "severity": issue.severity,
                "code": issue.code,
                "dataset": issue.dataset,
                "column": issue.column,
                "row_count": issue.row_count,
                "message": issue.message,
            }
        )
    return pd.DataFrame(rows)


def _build_join_summary_sheet(result: AnalysisResult) -> pd.DataFrame:
    if not result.join_reports:
        return pd.DataFrame([{"join_name": "none", "warning": "No joins were applied."}])

    rows = []
    for report in result.join_reports:
        rows.append(
            {
                "join_name": report.join_name,
                "left_dataset": report.left_dataset,
                "right_dataset": report.right_dataset,
                "join_type": report.join_type,
                "join_key_left": report.join_key_left,
                "join_key_right": report.join_key_right,
                "left_rows_before": report.left_rows_before,
                "left_rows_after": report.left_rows_after,
                "matched_rows": report.matched_rows,
                "unmatched_rows": report.unmatched_rows,
                "match_rate": report.match_rate,
                "duplicate_key_count_right": report.duplicate_key_count_right,
                "row_inflation": report.row_inflation,
                "warnings": " | ".join(report.warnings) if report.warnings else "",
            }
        )
    return pd.DataFrame(rows)


def _build_recommendations_sheet(result: AnalysisResult) -> pd.DataFrame:
    if not result.recommendations:
        return pd.DataFrame({"Recommendation": ["No recommendations available"]})
    return pd.DataFrame({"Recommendation": result.recommendations})


def _build_quality_issue_sheet(result: AnalysisResult) -> pd.DataFrame:
    if not result.quality_issues:
        return pd.DataFrame({"Issue": ["No major issues detected"]})
    return pd.DataFrame({"Issue": result.quality_issues})


def _build_alerts_sheet(result: AnalysisResult) -> pd.DataFrame:
    if result.alerts is None or result.alerts.empty:
        return pd.DataFrame({"alert_type": [], "details": []})
    return result.alerts.copy()


def _build_preflight_sheet(result: AnalysisResult) -> pd.DataFrame:
    preflight = evaluate_preflight(result)

    rows = [
        {"check": "overall_status", "status": preflight.overall_status, "details": "GREEN / AMBER / RED"},
        {"check": "risk_score", "status": preflight.overall_status, "details": preflight.risk_score},
        {
            "check": "validation_errors",
            "status": "RED" if preflight.metrics["validation_errors"] > 0 else "GREEN",
            "details": preflight.metrics["validation_errors"],
        },
        {
            "check": "validation_warnings",
            "status": "AMBER" if preflight.metrics["validation_warnings"] > 0 else "GREEN",
            "details": preflight.metrics["validation_warnings"],
        },
        {
            "check": "validation_info_count",
            "status": "INFO",
            "details": preflight.metrics["validation_infos"],
        },
        {
            "check": "join_reports_present",
            "status": "INFO" if preflight.metrics["total_join_reports"] == 0 else "GREEN",
            "details": preflight.metrics["total_join_reports"],
        },
        {
            "check": "joins_below_warning_match",
            "status": "AMBER" if preflight.metrics["joins_below_warning_match"] > 0 else "GREEN",
            "details": preflight.metrics["joins_below_warning_match"],
        },
        {
            "check": "joins_below_fail_match",
            "status": "RED" if preflight.metrics["joins_below_fail_match"] > 0 else "GREEN",
            "details": preflight.metrics["joins_below_fail_match"],
        },
        {
            "check": "joins_with_duplicate_keys",
            "status": "AMBER" if preflight.metrics["joins_with_duplicate_keys"] > 0 else "GREEN",
            "details": preflight.metrics["joins_with_duplicate_keys"],
        },
        {
            "check": "joins_with_severe_duplicate_keys",
            "status": "RED" if preflight.metrics["joins_with_severe_duplicate_keys"] > 0 else "GREEN",
            "details": preflight.metrics["joins_with_severe_duplicate_keys"],
        },
        {
            "check": "joins_with_row_inflation",
            "status": "AMBER" if preflight.metrics["joins_with_row_inflation"] > 0 else "GREEN",
            "details": preflight.metrics["joins_with_row_inflation"],
        },
        {
            "check": "joins_with_severe_row_inflation",
            "status": "RED" if preflight.metrics["joins_with_severe_row_inflation"] > 0 else "GREEN",
            "details": preflight.metrics["joins_with_severe_row_inflation"],
        },
        {
            "check": "join_warning_count",
            "status": "AMBER" if preflight.metrics["join_warning_count"] > 0 else "GREEN",
            "details": preflight.metrics["join_warning_count"],
        },
    ]

    if preflight.reasons:
        rows.append(
            {
                "check": "risk_reasons",
                "status": preflight.overall_status,
                "details": " | ".join(preflight.reasons[:10]),
            }
        )

    return pd.DataFrame(rows)


def _format_preflight_sheet(writer: pd.ExcelWriter) -> None:
    # workbook = writer.book
    workbook = cast(Workbook, writer.book)
    sheet_name = _safe_sheet_name("Preflight_Status")
    if sheet_name not in workbook.sheetnames:
        return

    ws = workbook[sheet_name]

    green_fill = PatternFill(fill_type="solid", fgColor="C6EFCE")
    amber_fill = PatternFill(fill_type="solid", fgColor="FFEB9C")
    red_fill = PatternFill(fill_type="solid", fgColor="FFC7CE")

    status_col = None
    for idx, cell in enumerate(ws[1], start=1):
        if cell.value == "status":
            status_col = idx
            break

    if status_col is None:
        return

    for row in range(2, ws.max_row + 1):
        cell = ws.cell(row=row, column=status_col)
        value = str(cell.value).upper() if cell.value is not None else ""

        if value == "GREEN":
            cell.fill = green_fill
        elif value == "AMBER":
            cell.fill = amber_fill
        elif value == "RED":
            cell.fill = red_fill



def _write_validation_samples(writer: pd.ExcelWriter, result: AnalysisResult) -> None:
    if result.validation_report is None:
        return

    sample_index = 1
    for issue in result.validation_report.issues:
        if issue.sample_rows is not None and not issue.sample_rows.empty:
            _write_df(writer, issue.sample_rows, f"ValidationSample_{sample_index}")
            sample_index += 1


def _write_unmatched_samples(writer: pd.ExcelWriter, result: AnalysisResult) -> None:
    if not result.join_reports:
        return

    sample_index = 1
    for report in result.join_reports:
        if report.unmatched_sample is not None and not report.unmatched_sample.empty:
            _write_df(writer, report.unmatched_sample, f"Unmatched_{sample_index}")
            sample_index += 1


def _format_worksheet(ws) -> None:
    ws.freeze_panes = "A2"

    header_font = Font(bold=True)
    for cell in ws[1]:
        cell.font = header_font

    for column_cells in ws.columns:
        max_length = 0
        col_idx = column_cells[0].column
        column_letter = get_column_letter(col_idx)

        for cell in column_cells:
            value = "" if cell.value is None else str(cell.value)
            if len(value) > max_length:
                max_length = len(value)

        adjusted_width = min(max(max_length + 2, 12), 40)
        ws.column_dimensions[column_letter].width = adjusted_width


def _apply_numeric_formats(writer: pd.ExcelWriter) -> None:
    # workbook = writer.book
    workbook = cast(Workbook, writer.book)
    for ws in workbook.worksheets:
        header_map = {cell.value: idx + 1 for idx, cell in enumerate(ws[1]) if cell.value}

        if "match_rate" in header_map:
            col_letter = get_column_letter(header_map["match_rate"])
            for row in range(2, ws.max_row + 1):
                ws[f"{col_letter}{row}"].number_format = "0.00%"

        numeric_like_headers = {
            "row_count",
            "left_rows_before",
            "left_rows_after",
            "matched_rows",
            "unmatched_rows",
            "duplicate_key_count_right",
            "row_inflation",
        }

        for header in numeric_like_headers:
            if header in header_map:
                col_letter = get_column_letter(header_map[header])
                for row in range(2, ws.max_row + 1):
                    ws[f"{col_letter}{row}"].number_format = "#,##0"


def _format_kpi_sheet(writer: pd.ExcelWriter) -> None:
    # workbook = writer.book
    workbook = cast(Workbook, writer.book)
    sheet_name = _safe_sheet_name("KPI_Summary")
    if sheet_name not in workbook.sheetnames:
        return

    ws = workbook[sheet_name]
    for row in range(2, ws.max_row + 1):
        ws[f"A{row}"].font = Font(bold=True)


def build_excel_report(result: AnalysisResult) -> bytes:
    output = io.BytesIO()

    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        _write_df(writer, _build_preflight_sheet(result), "Preflight_Status")
        _write_df(writer, _build_metadata_sheet(result), "Metadata")
        _write_df(writer, result.cleaned_df, "Cleaned_Data")
        _write_df(writer, _build_kpi_sheet(result), "KPI_Summary")
        _write_df(writer, _build_quality_issue_sheet(result), "Quality_Issues")
        _write_df(writer, _build_validation_sheet(result), "Validation_Report")
        _write_df(writer, _build_join_summary_sheet(result), "Join_Summary")
        _write_df(writer, _build_alerts_sheet(result), "Alerts")
        _write_df(writer, _build_recommendations_sheet(result), "Recommendations")

        for sheet_name, chart_df in result.chart_data.items():
            if isinstance(chart_df, pd.DataFrame):
                _write_df(writer, chart_df, sheet_name)

        _write_validation_samples(writer, result)
        _write_unmatched_samples(writer, result)

        workbook = writer.book
        for ws in workbook.worksheets:
            _format_worksheet(ws)

        _apply_numeric_formats(writer)
        _format_kpi_sheet(writer)
        _format_preflight_sheet(writer)

    output.seek(0)
    return output.getvalue()