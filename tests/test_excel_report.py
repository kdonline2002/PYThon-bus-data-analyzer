from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from reporting.excel_report import build_excel_report
from utils.models import AnalysisResult, ValidationReport


def test_excel_report_generates_workbook():
    validation_report = ValidationReport(dataset="sales")

    result = AnalysisResult(
        dataset_type="sales",
        cleaned_df=pd.DataFrame({
            "date": ["2025-01-01"],
            "product": ["A"],
            "quantity": [1],
            "sales_amount": [100],
        }),
        quality_issues=[],
        kpis={"Total Sales": "100.00"},
        alerts=pd.DataFrame(columns=["alert_type", "details"]),
        recommendations=["Keep monitoring sales."],
        chart_data={"top_products": pd.DataFrame({"product": ["A"], "sales_amount": [100]})},
        validation_report=validation_report,
        join_reports=[],
    )

    excel_bytes = build_excel_report(result)

    assert isinstance(excel_bytes, bytes)
    assert len(excel_bytes) > 0

    workbook = load_workbook(BytesIO(excel_bytes))

    assert "Preflight_Status" in workbook.sheetnames
    assert "Metadata" in workbook.sheetnames
    assert "Cleaned_Data" in workbook.sheetnames
    assert "KPI_Summary" in workbook.sheetnames
    assert "Validation_Report" in workbook.sheetnames


def test_excel_report_handles_validation_issues():
    validation_report = ValidationReport(dataset="sales")
    validation_report.add_issue(
        code="missing_required_field",
        severity="error",
        message="Missing quantity",
    )

    result = AnalysisResult(
        dataset_type="sales",
        cleaned_df=pd.DataFrame({"product": ["A"]}),
        quality_issues=["Missing quantity"],
        kpis={},
        alerts=pd.DataFrame(columns=["alert_type", "details"]),
        recommendations=[],
        chart_data={},
        validation_report=validation_report,
        join_reports=[],
    )

    excel_bytes = build_excel_report(result)
    workbook = load_workbook(BytesIO(excel_bytes))

    assert "Validation_Report" in workbook.sheetnames